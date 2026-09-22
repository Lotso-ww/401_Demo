from pathlib import Path
from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DELIVERABLES = ROOT / "docs" / "deliverables"


def escape_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", "<br>").strip()


def block_items(document):
    paragraphs = {id(p._p): p for p in document.paragraphs}
    tables = {id(t._tbl): t for t in document.tables}
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p") and id(child) in paragraphs:
            yield "paragraph", paragraphs[id(child)]
        elif child.tag == qn("w:tbl") and id(child) in tables:
            yield "table", tables[id(child)]


def paragraph_markdown(paragraph):
    text = paragraph.text.strip()
    if not text:
        return ""
    style = paragraph.style.name if paragraph.style else ""
    if style.startswith("Heading 1"):
        return f"# {text}"
    if style.startswith("Heading 2"):
        return f"## {text}"
    if style.startswith("Heading 3"):
        return f"### {text}"
    if "List Bullet" in style:
        return f"- {text}"
    return text


def table_markdown(table):
    rows = []
    for row in table.rows:
        rows.append([escape_cell(cell.text) for cell in row.cells])
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = "| " + " | ".join(rows[0]) + " |"
    divider = "| " + " | ".join("---" for _ in range(width)) + " |"
    body = ["| " + " | ".join(row) + " |" for row in rows[1:]]
    return "\n".join([header, divider, *body])


def convert(source: Path, target: Path):
    document = Document(source)
    blocks = []
    for kind, value in block_items(document):
        rendered = paragraph_markdown(value) if kind == "paragraph" else table_markdown(value)
        if rendered:
            blocks.append(rendered)

    text = "\n\n".join(blocks).strip() + "\n"
    text = text.replace("实际结果：待执行<br>结论：待执行", "实际结果：待执行<br>结论：待执行")
    target.write_text(text, encoding="utf-8", newline="\n")
    print(f"created {target} ({len(text.splitlines())} lines)")


def main():
    convert(
        DELIVERABLES / "TLS401_RFID_CCD_联动演示项目_软件概要设计说明书_V1.0.docx",
        DELIVERABLES / "TLS401_RFID_CCD_联动演示项目_软件概要设计说明书_V1.0.md",
    )
    convert(
        DELIVERABLES / "TLS401_RFID_CCD_联动演示项目_系统测试用例_V1.0.docx",
        DELIVERABLES / "TLS401_RFID_CCD_联动演示项目_系统测试用例_V1.0.md",
    )


if __name__ == "__main__":
    main()
