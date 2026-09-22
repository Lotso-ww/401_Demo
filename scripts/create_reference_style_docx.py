from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "docs" / "deliverables" / "reference-style"


def configure_run(run, size=10.5, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def configure_document(document, landscape=False):
    section = document.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.left_margin = Cm(1.25)
        section.right_margin = Cm(1.25)
    else:
        section.left_margin = Cm(2.45)
        section.right_margin = Cm(2.1)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(4)
    for style_name, size, color in (("Heading 1", 16, (31, 78, 121)), ("Heading 2", 13, (47, 84, 150)), ("Heading 3", 11, (31, 78, 121))):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    configure_run(header.add_run("TLS401 RFID + CCD 联动演示项目"), 8.5, color=(100, 100, 100))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    configure_run(footer.add_run("第 "), 8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    configure_run(footer.add_run(" 页"), 8)


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    props.append(element)


def keep_row(row):
    props = row._tr.get_or_add_trPr()
    element = OxmlElement("w:cantSplit")
    props.append(element)


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    props.append(element)


def table_cell(cell, value, size=8.8, header=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for line_index, line in enumerate(value.replace("<br>", "\n").split("\n")):
        if line_index:
            paragraph.add_run().add_break()
        configure_run(paragraph.add_run(line), size, bold=header, color=(255, 255, 255) if header else None)
    if header:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].startswith("|"):
        parts = [part.strip() for part in lines[index].strip().strip("|").split("|")]
        if not all(part.replace("-", "").replace(":", "") == "" for part in parts):
            rows.append(parts)
        index += 1
    return rows, index


def add_table(document, rows):
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for col, value in enumerate(rows[0]):
        shade(table.rows[0].cells[col], "4472C4")
        table_cell(table.rows[0].cells[col], value, 8.3, header=True)
    repeat_header(table.rows[0])
    for values in rows[1:]:
        row = table.add_row()
        keep_row(row)
        for col in range(columns):
            table_cell(row.cells[col], values[col] if col < len(values) else "", 7.4 if columns >= 8 else 8.8)
    document.add_paragraph()


def add_paragraph(document, text, bullet=False, code=False):
    if bullet:
        paragraph = document.add_paragraph(style="List Bullet")
    else:
        paragraph = document.add_paragraph()
    if not bullet and not code:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing = 1.3
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    configure_run(run, 8.5 if code else 10.5)
    if code:
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")


def make_docx(source, target, landscape=False):
    document = Document()
    configure_document(document, landscape)
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    first_title = True
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            add_paragraph(document, line, code=True)
        elif line.startswith("| "):
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        elif line.startswith("# "):
            text = line[2:].strip()
            if first_title:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(18)
                paragraph.paragraph_format.space_after = Pt(18)
                configure_run(paragraph.add_run(text), 22, bold=True, color=(31, 78, 121))
                first_title = False
            else:
                document.add_heading(text, level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
        elif line.startswith("- "):
            add_paragraph(document, line[2:].strip(), bullet=True)
        elif stripped:
            add_paragraph(document, stripped)
        index += 1
    document.save(target)
    print(target)


def main():
    make_docx(
        SOURCE_DIR / "TLS401_RFID_CCD_联动演示项目_软件概要设计说明书_规范学习版_V1.0.md",
        SOURCE_DIR / "TLS401_RFID_CCD_联动演示项目_软件概要设计说明书_规范学习版_V1.0.docx",
    )
    make_docx(
        SOURCE_DIR / "TLS401_RFID_CCD_联动演示项目_系统测试用例_规范学习版_V1.0.md",
        SOURCE_DIR / "TLS401_RFID_CCD_联动演示项目_系统测试用例_规范学习版_V1.0.docx",
        landscape=True,
    )


if __name__ == "__main__":
    main()
